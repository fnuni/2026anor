"""Recompute the anonymous expert-walkthrough summary from the released workbook.

The workbook contains the checked Word-control tags in its ``Raw`` sheet and
the pre-specified case, option, selector, repetition and answer-key tables in
the remaining sheets. Summary values are independently recomputed from inputs;
cached formula results are checked separately for errors and flag consistency.
"""
from __future__ import annotations

import json
import statistics
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from scipy.stats import binom


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "expert_consultation" / "Expert_consultation_data.xlsx"
EXPECTED = ROOT / "data" / "expert_walkthrough_summary.json"


def rows(ws):
    return list(ws.iter_rows(values_only=True))


def one(tags, prefix):
    """Return the one-based position of the single checked tag for prefix."""
    prefix = prefix + "_"
    values = sorted(int(tag.rsplit("_", 1)[1]) for tag in tags if tag.startswith(prefix))
    if len(values) > 1:
        raise AssertionError(f"Several answers for {prefix[:-1]}: {values}")
    return values[0] if values else None


def mean(values):
    return round(statistics.mean(values), 2)


def compute():
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)

    raw = defaultdict(set)
    for code, tag, value in rows(wb["Raw"])[1:]:
        if code and tag and value == 1:
            raw[code].add(tag)

    response_rows = rows(wb["Responses"])
    response_header = response_rows[1]
    code_col = response_header.index("Code")
    form_col = response_header.index("Form")
    forms = {
        row[code_col]: row[form_col]
        for row in response_rows[2:]
        if row[code_col] and row[code_col] != "EXAMPLE"
    }

    sequence = defaultdict(list)
    for form, case, instance, condition, r1, r2, r3, _note in rows(wb["Sequence"])[1:]:
        sequence[form].append((int(case), int(instance), condition, (int(r1), int(r2), int(r3))))
    for form in sequence:
        sequence[form].sort()

    options = defaultdict(dict)
    portfolio_labels = defaultdict(list)
    for key, instance, label, n1, n2, n3, _k, in_portfolio, _reachable, *_ in rows(wb["Data"])[1:]:
        if not key:
            continue
        instance = int(instance)
        options[instance][label] = (float(n1), float(n2), float(n3))
        if in_portfolio == 1:
            portfolio_labels[instance].append(label)

    selector = {}
    for _key, instance, theta, label in rows(wb["Selector"])[1:]:
        selector[(int(instance), round(float(theta), 1))] = label

    repetition = {}
    repeat_letters = defaultdict(list)
    for _key, form, shown, original in rows(wb["Repetition"])[1:]:
        repetition[(form, shown)] = original
        if shown != "N":
            repeat_letters[form].append(shown)

    vignette_keys = {}
    for form, v11, v12, v21, v22 in rows(wb["VignetteKeys"])[1:]:
        vignette_keys[form] = ((int(v11), int(v12)), (int(v21), int(v22)))

    def regret(instance, label, weights):
        def utility(norm):
            return -sum(w * x for w, x in zip(weights, norm)) / sum(weights)

        best = max(utility(options[instance][lab]) for lab in portfolio_labels[instance])
        return best - utility(options[instance][label])

    quiz_keys = (2, 1, 2, 2, 2)
    quiz_correct = 0
    reading_correct = 0
    portfolio = []
    random_zero_probs = []
    shown_option_counts = []
    nominal_regrets = []
    dial_regrets = []
    single = []
    repeat_same = 0
    confidence = {"A": [], "B": []}
    format_a_risk = []
    decisive = []
    experience_over_five = 0
    vignette = {"dial": {"correct": 0, "ease": []}, "summary": {"correct": 0, "ease": []}}
    sus = []
    agreement = {"understood_dial": [], "options_helped": [], "prefer_single": []}
    requests = {"workloads": 0, "past_days": 0, "written_explanation": 0}

    for code in sorted(forms):
        tags = raw[code]
        form = forms[code]
        if "ADESIONE_1" not in tags:
            raise AssertionError(f"Missing consent: {code}")
        experience_over_five += "P0B_2" in tags
        weights = tuple(sum(tag.startswith(f"TOK_{key}_") for tag in tags) for key in ("C", "T", "S"))
        if sum(weights) != 10:
            raise AssertionError(f"Token total differs from 10: {code}, {weights}")
        quiz_correct += sum(one(tags, f"Q{i}") == key for i, key in enumerate(quiz_keys, 1))
        original_choices = {}

        for case, instance, condition, reading_keys in sequence[form]:
            reading_correct += sum(one(tags, f"C{case}_L{i}") == key for i, key in enumerate(reading_keys, 1))
            d1 = one(tags, f"C{case}_D1")
            conf = one(tags, f"C{case}_D5")
            if condition == "A":
                single.append(d1)
                confidence["A"].append(conf)
                risk = one(tags, f"C{case}_D2")
                if risk is not None:
                    format_a_risk.append(risk)
                continue

            confidence["B"].append(conf)
            if condition == "R":
                shown_labels = ["N"] + repeat_letters[form]
                shown = shown_labels[d1 - 1]
                selected = repetition[(form, shown)]
                evaluated_labels = ["N"] + [repetition[(form, label)]
                                               for label in repeat_letters[form]]
                repeat_same += selected == original_choices[instance]
            else:
                shown_labels = ["N"] + portfolio_labels[instance]
                evaluated_labels = shown_labels
                selected = shown_labels[d1 - 1]
                original_choices[instance] = selected

            portfolio.append((selected, regret(instance, selected, weights)))
            shown_option_counts.append(len(shown_labels))
            random_zero_probs.append(
                sum(regret(instance, label, weights) <= 0.0005
                    for label in evaluated_labels) / len(evaluated_labels)
            )
            nominal_regrets.append(regret(instance, "N", weights))
            d4 = one(tags, f"C{case}_D4")
            if d4 is not None:
                decisive.append(d4)

            d2 = one(tags, f"C{case}_D2")
            if d2 is not None and d2 <= 9:
                theta = round(d2 / 10, 1)
                chosen_regret = regret(instance, selector[(instance, theta)], weights)
                attainable = [
                    regret(instance, selector[(instance, round(i / 10, 1))], weights)
                    for i in range(1, 10)
                ]
                dial_regrets.append((chosen_regret, min(attainable)))

        for i, presentation in ((1, "dial"), (2, "summary")):
            vignette[presentation]["correct"] += sum(
                one(tags, f"V{i}_{j}") == key for j, key in enumerate(vignette_keys[form][i - 1], 1)
            )
            vignette[presentation]["ease"].append(one(tags, f"V{i}_3"))

        answers = [one(tags, f"S{i}") for i in range(1, 11)]
        if all(value is not None for value in answers):
            score = sum(answers[i] - 1 for i in (0, 2, 4, 6, 8))
            score += sum(5 - answers[i] for i in (1, 3, 5, 7, 9))
            sus.append(score * 2.5)

        for key, field in (("G1", "understood_dial"), ("G3", "options_helped"), ("G6", "prefer_single")):
            agreement[field].append(one(tags, key))
        for position, field in ((1, "workloads"), (3, "past_days"), (4, "written_explanation")):
            requests[field] += f"F2_{position}" in tags

    random_expected = sum(random_zero_probs)
    random_p = float(binom.sf(
        sum(value <= 0.0005 for _, value in portfolio) - 1,
        len(portfolio), random_expected / len(portfolio)))
    return {
        "booklets": len(forms),
        "experience_over_five": experience_over_five,
        "comprehension_correct": quiz_correct,
        "reading_correct": reading_correct,
        "portfolio_decisions": len(portfolio),
        "portfolio_selected_N": sum(label == "N" for label, _ in portfolio),
        "portfolio_zero_regret": sum(value <= 0.0005 for _, value in portfolio),
        "portfolio_mean_regret": mean(value for _, value in portfolio),
        "nominal_mean_regret": mean(nominal_regrets),
        "random_choice_mean_options": round(statistics.mean(shown_option_counts), 1),
        "random_choice_zero_regret_expected": round(random_expected, 1),
        "random_choice_binomial_p": random_p,
        "dial_set": len(dial_regrets),
        "dial_within_0_05": sum(value <= best + 0.05 + 1e-12 for value, best in dial_regrets),
        "repeated_same_option": repeat_same,
        "single_recommendation": {
            "adopt": single.count(1),
            "alternatives": single.count(2),
            "reject": single.count(3),
        },
        "format_A_risk_median": statistics.median(format_a_risk),
        "decisive_column": {
            "one_in_twenty": decisive.count(5),
            "selector": decisive.count(10),
            "average_cost": decisive.count(3),
            "answered": len(decisive),
        },
        "mean_confidence": {"B": mean(confidence["B"]), "A": mean(confidence["A"])},
        "vignettes": {
            key: {"correct": value["correct"], "ease": mean(value["ease"])}
            for key, value in vignette.items()
        },
        "SUS": {"n": len(sus), "mean": mean(sus), "range": [min(sus), max(sus)]},
        "agreement": {key: mean(value) for key, value in agreement.items()},
        "requested_features": requests,
    }


def compare(actual, expected, path="summary"):
    if isinstance(expected, dict):
        if set(actual) != set(expected):
            raise AssertionError(f"Keys differ at {path}: {set(actual) ^ set(expected)}")
        for key in expected:
            compare(actual[key], expected[key], f"{path}.{key}")
    elif isinstance(expected, list):
        if actual != expected:
            raise AssertionError(f"Value differs at {path}: {actual!r} != {expected!r}")
    elif actual != expected:
        raise AssertionError(f"Value differs at {path}: {actual!r} != {expected!r}")


def verify_cached_checks():
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    errors = [(ws.title, cell.coordinate, cell.value) for ws in wb
              for row in ws for cell in row if cell.data_type == "e"]
    assert not errors, errors
    checks = rows(wb["Checks"])
    header = checks[0]
    total_column = header.index("Number of X flags")
    flag_columns = [i for i, value in enumerate(header)
                    if value and "flag" in value.lower() and i != total_column]
    counts = {}
    for row in checks[1:]:
        if row[0] and row[0].startswith("P"):
            count = sum(row[i] == "X" for i in flag_columns)
            assert row[total_column] == count, (row[0], row[total_column], count)
            counts[row[0]] = count
    assert len(counts) == 10 and sum(counts.values()) == 11, counts
    wb.close()
    return counts


if __name__ == "__main__":
    actual = compute()
    expected = json.loads(EXPECTED.read_text(encoding="utf-8"))
    compare(actual, expected)
    counts = verify_cached_checks()
    print(json.dumps(actual, indent=2, ensure_ascii=False))
    print("expert walkthrough: all released summary values verified")
    print("cached formulas: no errors; inconsistency flags:", counts)
